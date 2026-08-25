"""Smoke-test the real OCR and legacy DOC tools without touching project data."""

from __future__ import annotations

import json
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

from onetech_ao.config import Settings
from onetech_ao.documents import extract_text


def run(command: list[str], timeout: int = 120) -> None:
    completed = subprocess.run(
        command,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        check=False,
        shell=False,
        timeout=timeout,
    )
    if completed.returncode != 0:
        detail = " ".join((completed.stderr or completed.stdout).split())[:300]
        raise RuntimeError(f"Command failed ({completed.returncode}): {detail}")


def convert(source: Path, output_dir: Path, target: str, profile: Path) -> Path:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice is unavailable")
    output_dir.mkdir(parents=True, exist_ok=True)
    profile.mkdir(parents=True, exist_ok=True)
    run(
        [
            executable,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            "--convert-to",
            target,
            "--outdir",
            str(output_dir),
            str(source),
        ]
    )
    extension = target.split(":", 1)[0]
    output = output_dir / f"{source.stem}.{extension}"
    if not output.is_file():
        raise RuntimeError(f"LibreOffice did not produce {output.name}")
    return output


def main() -> int:
    phrase = "ONETECH OCR VALIDATION 2026"
    with tempfile.TemporaryDirectory(prefix="onetech-tools-smoke-") as raw:
        root = Path(raw)
        source_docx = root / "source.docx"
        document = Document()
        paragraph = document.add_paragraph()
        run_text = paragraph.add_run(phrase)
        run_text.bold = True
        run_text.font.size = Pt(26)
        document.save(source_docx)

        source_pdf = convert(
            source_docx,
            root / "source-pdf",
            "pdf",
            root / "profile-pdf",
        )
        legacy_doc = convert(
            source_docx,
            root / "source-doc",
            "doc:MS Word 97",
            root / "profile-doc",
        )

        pdftoppm = shutil.which("pdftoppm")
        if not pdftoppm:
            raise RuntimeError("pdftoppm is unavailable")
        image_prefix = root / "scan-page"
        run(
            [
                pdftoppm,
                "-f",
                "1",
                "-l",
                "1",
                "-singlefile",
                "-r",
                "200",
                "-png",
                str(source_pdf),
                str(image_prefix),
            ]
        )
        image = image_prefix.with_suffix(".png")
        scanned_docx = root / "scanned.docx"
        scanned_document = Document()
        scanned_document.add_picture(str(image), width=Inches(7))
        scanned_document.save(scanned_docx)
        scanned_pdf = convert(
            scanned_docx,
            root / "scanned-pdf",
            "pdf",
            root / "profile-scanned",
        )

        settings = Settings(
            postgres_db="smoke",
            postgres_user="smoke",
            postgres_password="smoke",
            project_root=root,
            document_root=root / "documents",
            document_ocr_languages="eng",
            document_ocr_max_pages=2,
        )
        ocr_text = extract_text(scanned_pdf, "pdf", settings) or ""
        doc_text = extract_text(legacy_doc, "doc", settings) or ""
        normalized_ocr = " ".join(ocr_text.upper().split())
        normalized_doc = " ".join(doc_text.upper().split())
        if phrase not in normalized_ocr:
            raise RuntimeError(
                "OCR output did not contain the validation phrase: "
                + repr(ocr_text[:300])
            )
        if phrase not in normalized_doc:
            raise RuntimeError("DOC conversion output did not contain the validation phrase")

        print(
            json.dumps(
                {
                    "status": "ok",
                    "ocr": True,
                    "legacy_doc": True,
                    "temporary_files_removed_on_exit": True,
                },
                ensure_ascii=False,
            )
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
