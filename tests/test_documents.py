from __future__ import annotations

from io import BytesIO
from pathlib import Path
import subprocess
from typing import Any

import pytest
from docx import Document
from pypdf import PdfWriter

from onetech_ao.config import Settings
from onetech_ao import documents as documents_module
from onetech_ao.documents import (
    DocumentProcessingError,
    DocumentProcessor,
    _detect_document_type,
    _validate_rendered_image,
    extract_text,
    validate_public_document_url,
)


class FakeResponse:
    def __init__(self, body: bytes, url: str, content_type: str) -> None:
        self.body = body
        self.url = url
        self.headers = {
            "Content-Length": str(len(body)),
            "Content-Type": content_type,
        }

    def raise_for_status(self) -> None:
        return None

    def iter_content(self, chunk_size: int) -> list[bytes]:
        return [
            self.body[index : index + chunk_size]
            for index in range(0, len(self.body), chunk_size)
        ]


class FakeSession:
    def __init__(self, response: FakeResponse) -> None:
        self.response = response
        self.headers: dict[str, str] = {}
        self.calls = 0
        self.closed = False

    def mount(self, prefix: str, adapter: object) -> None:
        return None

    def get(self, url: str, **kwargs: object) -> FakeResponse:
        self.calls += 1
        return self.response

    def close(self) -> None:
        self.closed = True


class FakeDatabase:
    def __init__(self, task: dict[str, Any]) -> None:
        self.task = task
        self.saved: dict[str, Any] | None = None
        self.failed: tuple[int, str] | None = None

    def check_connection(self) -> dict[str, str]:
        return {"database": "test"}

    def documents_to_process(
        self,
        limit: int,
        retry_errors: bool,
        reprocess_extracted: bool = False,
        tender_id: str | None = None,
    ) -> list[dict[str, Any]]:
        return [self.task]

    def save_document_processing(self, **kwargs: Any) -> None:
        self.saved = kwargs

    def fail_document_processing(
        self, document_id: int, error: str, **kwargs: Any
    ) -> None:
        self.failed = (document_id, error)


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        postgres_db="test",
        postgres_user="test",
        postgres_password="secret",
        project_root=tmp_path,
        document_root=tmp_path / "data" / "documents",
    )


def _docx_bytes() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_heading("Cahier des charges", level=1)
    document.add_paragraph("Date limite : 30 septembre 2026 à 10h00")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Caution"
    table.rows[0].cells[1].text = "5 000 TND"
    document.save(output)
    return output.getvalue()


def _blank_pdf_bytes() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(output)
    return output.getvalue()


def _install_fake_libreoffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        documents_module,
        "_find_executable",
        lambda configured, fallbacks: "libreoffice",
    )

    def fake_run(
        command: list[str], *, timeout_seconds: int, error_label: str
    ) -> subprocess.CompletedProcess[str]:
        output_dir = Path(command[command.index("--outdir") + 1])
        (output_dir / "source.docx").write_bytes(_docx_bytes())
        return subprocess.CompletedProcess(command, 0, stdout="convert ok", stderr="")

    monkeypatch.setattr(documents_module, "_run_command", fake_run)


def test_processor_downloads_hashes_and_extracts_docx(tmp_path: Path) -> None:
    url = "https://www.marchespublics.gov.tn/storage/tender/test/cdc.docx"
    task = {
        "id": 42,
        "tender_id": "Tender-42",
        "url": url,
        "type_document": "docx",
        "chemin_local": None,
        "empreinte_sha256": None,
    }
    database = FakeDatabase(task)
    session = FakeSession(
        FakeResponse(
            body=_docx_bytes(),
            url=url,
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )
    )

    result = DocumentProcessor(
        settings=_settings(tmp_path),
        database=database,  # type: ignore[arg-type]
        session=session,  # type: ignore[arg-type]
    ).process(limit=1)

    assert result.extraits == 1
    assert result.echecs == 0
    assert database.failed is None
    assert database.saved is not None
    assert database.saved["status"] == "extrait"
    assert database.saved["detected_type"] == "docx"
    assert len(database.saved["sha256"]) == 64
    assert "Date limite" in database.saved["extracted_text"]
    assert "5 000 TND" in database.saved["extracted_text"]
    assert (tmp_path / database.saved["relative_path"]).is_file()
    assert session.closed is True


def test_processor_rejects_non_haicop_url_without_downloading(tmp_path: Path) -> None:
    url = "https://example.com/storage/tender/evil.pdf"
    task = {
        "id": 7,
        "tender_id": "Tender-7",
        "url": url,
        "type_document": "pdf",
        "chemin_local": None,
        "empreinte_sha256": None,
    }
    database = FakeDatabase(task)
    session = FakeSession(FakeResponse(b"%PDF-test", url, "application/pdf"))

    result = DocumentProcessor(
        settings=_settings(tmp_path),
        database=database,  # type: ignore[arg-type]
        session=session,  # type: ignore[arg-type]
    ).process(limit=1)

    assert result.echecs == 1
    assert database.saved is None
    assert database.failed is not None
    assert "hôte non autorisé" in database.failed[1]
    assert session.calls == 0


def test_processor_ocr_extracts_a_scanned_pdf(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    url = "https://www.marchespublics.gov.tn/storage/tender/test/scanned.pdf"
    task = {
        "id": 43,
        "tender_id": "Tender-43",
        "url": url,
        "type_document": "pdf",
        "chemin_local": None,
        "empreinte_sha256": None,
    }
    database = FakeDatabase(task)
    session = FakeSession(FakeResponse(_blank_pdf_bytes(), url, "application/pdf"))
    calls: list[list[int]] = []

    def fake_ocr(
        path: Path, page_numbers: list[int], settings: Settings
    ) -> dict[int, str]:
        calls.append(page_numbers)
        return {1: "Texte reconnu par OCR en francais"}

    monkeypatch.setattr(documents_module, "_ocr_pdf_pages", fake_ocr)

    result = DocumentProcessor(
        settings=_settings(tmp_path),
        database=database,  # type: ignore[arg-type]
        session=session,  # type: ignore[arg-type]
    ).process(limit=1)

    assert result.extraits == 1
    assert result.sans_texte == 0
    assert calls == [[1]]
    assert database.saved is not None
    assert database.saved["status"] == "extrait"
    assert database.saved["detected_type"] == "pdf"
    assert "--- PAGE 1 ---" in database.saved["extracted_text"]
    assert "Texte reconnu" in database.saved["extracted_text"]


def test_pdf_without_text_stays_empty_when_ocr_is_disabled(tmp_path: Path) -> None:
    path = tmp_path / "blank.pdf"
    path.write_bytes(_blank_pdf_bytes())
    settings = _settings(tmp_path)
    settings = Settings(
        postgres_db=settings.postgres_db,
        postgres_user=settings.postgres_user,
        postgres_password=settings.postgres_password,
        project_root=settings.project_root,
        document_root=settings.document_root,
        document_ocr_enabled=False,
    )

    assert extract_text(path, "pdf", settings) is None


def test_processor_converts_and_extracts_legacy_doc(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _install_fake_libreoffice(monkeypatch)
    url = "https://www.marchespublics.gov.tn/storage/tender/test/legacy.doc"
    task = {
        "id": 44,
        "tender_id": "Tender-44",
        "url": url,
        "type_document": "doc",
        "chemin_local": None,
        "empreinte_sha256": None,
    }
    database = FakeDatabase(task)
    body = documents_module.OLE_SIGNATURE + (b"legacy-word-data" * 20)
    session = FakeSession(
        FakeResponse(body, url, "application/msword")
    )

    result = DocumentProcessor(
        settings=_settings(tmp_path),
        database=database,  # type: ignore[arg-type]
        session=session,  # type: ignore[arg-type]
    ).process(limit=1)

    assert result.extraits == 1
    assert result.non_supportes == 0
    assert database.saved is not None
    assert database.saved["status"] == "extrait"
    assert database.saved["detected_type"] == "doc"
    assert "Cahier des charges" in database.saved["extracted_text"]
    assert "5 000 TND" in database.saved["extracted_text"]


def test_processor_ocr_extracts_a_standalone_png(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    url = "https://www.marchespublics.gov.tn/storage/tender/test/notice.png"
    task = {
        "id": 45,
        "tender_id": "Tender-45",
        "url": url,
        "type_document": "png",
        "chemin_local": None,
        "empreinte_sha256": None,
    }
    database = FakeDatabase(task)
    body = documents_module.PNG_SIGNATURE + b"fake-image-data"
    session = FakeSession(FakeResponse(body, url, "image/png"))
    monkeypatch.setattr(
        documents_module,
        "_ocr_image_text",
        lambda path, settings: "Avis d'appel d'offres reconnu",
    )

    result = DocumentProcessor(
        settings=_settings(tmp_path),
        database=database,  # type: ignore[arg-type]
        session=session,  # type: ignore[arg-type]
    ).process(limit=1)

    assert result.extraits == 1
    assert result.non_supportes == 0
    assert database.saved is not None
    assert database.saved["detected_type"] == "png"
    assert database.saved["status"] == "extrait"


def test_ocr_reports_missing_system_tool(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        documents_module,
        "_find_executable",
        lambda configured, fallbacks: None,
    )

    with pytest.raises(DocumentProcessingError, match="OCR indisponible"):
        documents_module._ocr_pdf_pages(
            tmp_path / "scanned.pdf",
            [1],
            _settings(tmp_path),
        )


def test_ocr_uses_installed_language_subset_and_page_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pdf = tmp_path / "scanned.pdf"
    pdf.write_bytes(_blank_pdf_bytes())
    commands: list[list[str]] = []

    monkeypatch.setattr(
        documents_module,
        "_find_executable",
        lambda configured, fallbacks: configured,
    )

    def fake_run(
        command: list[str], *, timeout_seconds: int, error_label: str
    ) -> subprocess.CompletedProcess[str]:
        commands.append(command)
        if command[1:] == ["--list-langs"]:
            return subprocess.CompletedProcess(
                command,
                0,
                stdout="List of available languages in /tmp (2):\neng\nfra\n",
                stderr="",
            )
        if "pdftoppm" in command[0]:
            Path(command[-1]).with_suffix(".png").write_bytes(b"small-image")
            return subprocess.CompletedProcess(command, 0, stdout="", stderr="")
        return subprocess.CompletedProcess(
            command, 0, stdout="Texte OCR", stderr=""
        )

    monkeypatch.setattr(documents_module, "_run_command", fake_run)
    settings = Settings(
        postgres_db="test",
        postgres_user="test",
        postgres_password="secret",
        project_root=tmp_path,
        document_root=tmp_path / "data" / "documents",
        document_ocr_languages="fra+ara+eng",
        document_ocr_max_pages=1,
    )

    result = documents_module._ocr_pdf_pages(pdf, [1, 2], settings)

    assert result == {1: "Texte OCR"}
    tesseract_calls = [command for command in commands if "stdout" in command]
    assert len(tesseract_calls) == 1
    assert tesseract_calls[0][tesseract_calls[0].index("-l") + 1] == "fra+eng"


def test_ocr_refuses_excessive_temporary_render(tmp_path: Path) -> None:
    image = tmp_path / "render.png"
    image.write_bytes(b"too-large")
    settings = Settings(
        postgres_db="test",
        postgres_user="test",
        postgres_password="secret",
        project_root=tmp_path,
        document_root=tmp_path / "data" / "documents",
        document_ocr_max_render_bytes=len(b"too-large") - 1,
    )

    with pytest.raises(DocumentProcessingError, match="rendu temporaire trop volumineux"):
        _validate_rendered_image(image, 0, settings)


def test_ole_file_without_doc_extension_is_not_converted_as_word(tmp_path: Path) -> None:
    path = tmp_path / "legacy.xls.part"
    path.write_bytes(documents_module.OLE_SIGNATURE + b"not-a-word-document")

    assert _detect_document_type(path) == "ole"


def test_validate_public_document_url_accepts_only_public_storage() -> None:
    validate_public_document_url(
        "https://www.marchespublics.gov.tn/storage/tender/2026/08/test.pdf"
    )
