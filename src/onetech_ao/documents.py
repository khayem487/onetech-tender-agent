"""Téléchargement sûr et extraction locale des documents publics HAICOP."""

from __future__ import annotations

import hashlib
import logging
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import unquote, urlparse

import requests
from docx import Document
from pypdf import PdfReader
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from .config import Settings
from .database import Database


logging.getLogger("pypdf").setLevel(logging.ERROR)

ALLOWED_DOCUMENT_HOSTS = {"marchespublics.gov.tn", "www.marchespublics.gov.tn"}
SAFE_FILENAME_RE = re.compile(r"[^A-Za-z0-9._-]+")
PDF_SIGNATURE = b"%PDF-"
OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
PNG_SIGNATURE = bytes.fromhex("89504E470D0A1A0A")
JPEG_SIGNATURE = bytes.fromhex("FFD8FF")
SUPPORTED_DOCUMENT_TYPES = frozenset({"pdf", "docx", "doc", "png", "jpeg"})


class DocumentProcessingError(RuntimeError):
    """Erreur concise et sûre pendant le traitement d'un document."""


@dataclass(slots=True)
class DocumentTask:
    id: int
    tender_id: str
    url: str
    type_document: str | None
    chemin_local: str | None
    empreinte_sha256: str | None

    @classmethod
    def from_row(cls, row: dict[str, Any]) -> "DocumentTask":
        return cls(
            id=int(row["id"]),
            tender_id=str(row["tender_id"]),
            url=str(row["url"]),
            type_document=row.get("type_document"),
            chemin_local=row.get("chemin_local"),
            empreinte_sha256=row.get("empreinte_sha256"),
        )


@dataclass(slots=True)
class DownloadedDocument:
    relative_path: str
    absolute_path: Path
    sha256: str
    size_bytes: int
    mime_type: str | None
    detected_type: str


@dataclass(slots=True)
class DocumentProcessingResult:
    demandes: int = 0
    extraits: int = 0
    sans_texte: int = 0
    non_supportes: int = 0
    echecs: int = 0
    erreurs: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def validate_public_document_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme != "https":
        raise DocumentProcessingError("URL document refusée: HTTPS requis")
    if parsed.hostname not in ALLOWED_DOCUMENT_HOSTS:
        raise DocumentProcessingError("URL document refusée: hôte non autorisé")
    if not parsed.path.startswith("/storage/tender/"):
        raise DocumentProcessingError("URL document refusée: chemin HAICOP inattendu")


def _safe_filename(url: str, document_id: int) -> str:
    raw = unquote(PurePosixPath(urlparse(url).path).name)
    cleaned = SAFE_FILENAME_RE.sub("_", raw).strip("._")
    if not cleaned:
        cleaned = f"document-{document_id}"
    return f"{document_id}-{cleaned}"[:180]


def _detect_document_type(path: Path) -> str:
    with path.open("rb") as handle:
        signature = handle.read(8)
    if signature.startswith(PDF_SIGNATURE):
        return "pdf"
    if signature.startswith(OLE_SIGNATURE):
        # OLE is also used by old Excel/PowerPoint files.  Only send an
        # explicitly named legacy Word document to LibreOffice conversion.
        suffix = path.suffix.casefold()
        if suffix == ".part":
            suffix = path.with_suffix("").suffix.casefold()
        return "doc" if suffix == ".doc" else "ole"
    if signature.startswith(PNG_SIGNATURE):
        return "png"
    if signature.startswith(JPEG_SIGNATURE):
        return "jpeg"
    if zipfile.is_zipfile(path):
        try:
            with zipfile.ZipFile(path) as archive:
                names = set(archive.namelist())
            if "[Content_Types].xml" in names and "word/document.xml" in names:
                return "docx"
        except zipfile.BadZipFile as exc:
            raise DocumentProcessingError("archive DOCX invalide") from exc
    return "inconnu"


def _validate_docx_archive(path: Path, max_uncompressed_bytes: int) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000:
                raise DocumentProcessingError("DOCX refusé: trop de fichiers internes")
            total = sum(entry.file_size for entry in entries)
            if total > max_uncompressed_bytes:
                raise DocumentProcessingError("DOCX refusé: contenu décompressé trop volumineux")
    except zipfile.BadZipFile as exc:
        raise DocumentProcessingError("DOCX invalide") from exc


def _clean_text(parts: list[str], max_chars: int) -> str:
    lines: list[str] = []
    for part in parts:
        if not part or not part.strip():
            continue
        for line in part.splitlines():
            normalized = " ".join(line.split())
            if normalized:
                lines.append(normalized)
    return "\n".join(lines)[:max_chars]


def _find_executable(configured: str, fallbacks: tuple[str, ...]) -> str | None:
    """Resolve an explicitly configured command without invoking a shell."""

    candidate = Path(configured).expanduser()
    if candidate.is_absolute():
        return str(candidate) if candidate.is_file() else None
    for name in (configured, *fallbacks):
        resolved = shutil.which(name)
        if resolved:
            return resolved
    return None


def _run_command(
    command: list[str],
    *,
    timeout_seconds: int,
    error_label: str,
) -> subprocess.CompletedProcess[str]:
    """Run a bounded local converter and return UTF-8-safe output."""

    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            check=False,
            shell=False,
            timeout=timeout_seconds,
        )
    except subprocess.TimeoutExpired as exc:
        raise DocumentProcessingError(f"{error_label}: delai depasse") from exc
    except OSError as exc:
        raise DocumentProcessingError(
            f"{error_label}: executable indisponible"
        ) from exc
    if completed.returncode != 0:
        detail = " ".join((completed.stderr or completed.stdout or "").split())[:180]
        suffix = f" ({detail})" if detail else ""
        raise DocumentProcessingError(f"{error_label}: code {completed.returncode}{suffix}")
    return completed


def _available_tesseract_languages(
    executable: str, settings: Settings
) -> str:
    """Use available OCR languages rather than failing on one missing pack."""

    completed = _run_command(
        [executable, "--list-langs"],
        timeout_seconds=settings.document_ocr_timeout_seconds,
        error_label="verification OCR impossible",
    )
    available = {
        line.strip()
        for line in completed.stdout.splitlines()
        if re.fullmatch(r"[A-Za-z0-9_]+", line.strip())
    }
    requested = [
        language
        for language in settings.document_ocr_languages.split("+")
        if language in available
    ]
    if not requested:
        raise DocumentProcessingError(
            "OCR indisponible: aucune langue configuree n'est installee"
        )
    return "+".join(requested)


def _validate_rendered_image(
    image: Path,
    rendered_bytes: int,
    settings: Settings,
) -> int:
    """Bound temporary OCR disk usage before handing a bitmap to Tesseract."""

    try:
        size = image.stat().st_size
    except OSError as exc:
        raise DocumentProcessingError("rendu OCR impossible: image inaccessible") from exc
    total = rendered_bytes + size
    if total > settings.document_ocr_max_render_bytes:
        raise DocumentProcessingError("OCR refuse: rendu temporaire trop volumineux")
    return total


def _extract_docx_text(path: Path, settings: Settings) -> str | None:
    _validate_docx_archive(path, settings.document_max_uncompressed_bytes)
    try:
        document = Document(str(path))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
    except Exception as exc:
        raise DocumentProcessingError(
            f"extraction DOCX impossible: {type(exc).__name__}"
        ) from exc
    text = _clean_text(parts, settings.document_max_text_chars)
    return text or None


def _extract_legacy_doc_text(path: Path, settings: Settings) -> str | None:
    """Convert an OLE Word file in an isolated temporary profile, then parse it."""

    executable = _find_executable(
        settings.libreoffice_command,
        ("libreoffice", "soffice", "soffice.exe"),
    )
    if executable is None:
        raise DocumentProcessingError(
            "conversion DOC indisponible: LibreOffice est introuvable"
        )

    with tempfile.TemporaryDirectory(prefix="onetech-doc-") as temporary_raw:
        temporary = Path(temporary_raw)
        source = temporary / "source.doc"
        output_dir = temporary / "converted"
        profile_dir = temporary / "libreoffice-profile"
        output_dir.mkdir()
        profile_dir.mkdir()
        shutil.copyfile(path, source)
        _run_command(
            [
                executable,
                f"-env:UserInstallation={profile_dir.as_uri()}",
                "--headless",
                "--invisible",
                "--nologo",
                "--nodefault",
                "--nolockcheck",
                "--norestore",
                "--nofirststartwizard",
                "--convert-to",
                "docx:Office Open XML Text",
                "--outdir",
                str(output_dir),
                str(source),
            ],
            timeout_seconds=settings.document_conversion_timeout_seconds,
            error_label="conversion DOC impossible",
        )
        converted = output_dir / "source.docx"
        if not converted.is_file():
            candidates = list(output_dir.glob("*.docx"))
            if len(candidates) != 1:
                raise DocumentProcessingError(
                    "conversion DOC impossible: aucun DOCX produit"
                )
            converted = candidates[0]
        if converted.stat().st_size > settings.document_max_uncompressed_bytes:
            raise DocumentProcessingError("DOC converti refuse: fichier trop volumineux")
        return _extract_docx_text(converted, settings)


def _ocr_pdf_pages(
    path: Path,
    page_numbers: list[int],
    settings: Settings,
) -> dict[int, str]:
    """Render selected PDF pages and OCR them with bounded subprocesses."""

    pdftoppm = _find_executable(settings.pdftoppm_command, ("pdftoppm", "pdftoppm.exe"))
    tesseract = _find_executable(
        settings.tesseract_command,
        ("tesseract", "tesseract.exe"),
    )
    if pdftoppm is None or tesseract is None:
        missing = "pdftoppm" if pdftoppm is None else "tesseract"
        raise DocumentProcessingError(f"OCR indisponible: {missing} est introuvable")

    language_spec = _available_tesseract_languages(tesseract, settings)
    selected = page_numbers[: settings.document_ocr_max_pages]
    extracted: dict[int, str] = {}
    rendered_bytes = 0
    with tempfile.TemporaryDirectory(prefix="onetech-ocr-") as temporary_raw:
        temporary = Path(temporary_raw)
        for page_number in selected:
            prefix = temporary / f"page-{page_number}"
            _run_command(
                [
                    pdftoppm,
                    "-f",
                    str(page_number),
                    "-l",
                    str(page_number),
                    "-singlefile",
                    "-r",
                    str(settings.document_ocr_dpi),
                    "-png",
                    str(path),
                    str(prefix),
                ],
                timeout_seconds=settings.document_ocr_timeout_seconds,
                error_label=f"rendu OCR page {page_number} impossible",
            )
            image = prefix.with_suffix(".png")
            if not image.is_file():
                raise DocumentProcessingError(
                    f"rendu OCR page {page_number} impossible: image absente"
                )
            rendered_bytes = _validate_rendered_image(
                image, rendered_bytes, settings
            )
            completed = _run_command(
                [
                    tesseract,
                    str(image),
                    "stdout",
                    "-l",
                    language_spec,
                    "--psm",
                    "6",
                ],
                timeout_seconds=settings.document_ocr_timeout_seconds,
                error_label=f"OCR page {page_number} impossible",
            )
            page_text = _clean_text([completed.stdout], settings.document_max_text_chars)
            if page_text:
                extracted[page_number] = page_text
    return extracted


def _ocr_image_text(path: Path, settings: Settings) -> str | None:
    """OCR a standalone HAICOP PNG/JPEG under the same process bounds."""

    tesseract = _find_executable(
        settings.tesseract_command,
        ("tesseract", "tesseract.exe"),
    )
    if tesseract is None:
        raise DocumentProcessingError("OCR indisponible: tesseract est introuvable")
    language_spec = _available_tesseract_languages(tesseract, settings)
    completed = _run_command(
        [
            tesseract,
            str(path),
            "stdout",
            "-l",
            language_spec,
            "--psm",
            "6",
        ],
        timeout_seconds=settings.document_ocr_timeout_seconds,
        error_label="OCR image impossible",
    )
    text = _clean_text([completed.stdout], settings.document_max_text_chars)
    return text or None


def extract_text(path: Path, document_type: str, settings: Settings) -> str | None:
    if document_type == "pdf":
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted and reader.decrypt("") == 0:
                raise DocumentProcessingError("PDF chiffré sans mot de passe public")
            if len(reader.pages) > settings.document_max_pdf_pages:
                raise DocumentProcessingError("PDF refusé: trop de pages")
            page_texts: dict[int, str] = {}
            pages_without_text: list[int] = []
            for page_number, page in enumerate(reader.pages, start=1):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    page_texts[page_number] = page_text
                else:
                    pages_without_text.append(page_number)
        except DocumentProcessingError:
            raise
        except Exception as exc:
            raise DocumentProcessingError(f"extraction PDF impossible: {type(exc).__name__}") from exc
        if pages_without_text and settings.document_ocr_enabled:
            page_texts.update(_ocr_pdf_pages(path, pages_without_text, settings))
        parts: list[str] = []
        for page_number in sorted(page_texts):
            parts.append(f"--- PAGE {page_number} ---")
            parts.append(page_texts[page_number])
        text = _clean_text(parts, settings.document_max_text_chars)
        return text or None

    if document_type == "docx":
        return _extract_docx_text(path, settings)

    if document_type == "doc":
        return _extract_legacy_doc_text(path, settings)

    if document_type in {"png", "jpeg"}:
        if not settings.document_ocr_enabled:
            return None
        return _ocr_image_text(path, settings)

    return None


class DocumentProcessor:
    def __init__(
        self,
        settings: Settings,
        database: Database | None = None,
        session: requests.Session | None = None,
    ) -> None:
        self.settings = settings
        self.database = database or Database(settings)
        self.session = session or requests.Session()
        retry = Retry(
            total=3,
            connect=3,
            read=3,
            status=3,
            backoff_factor=0.5,
            status_forcelist=(429, 500, 502, 503, 504),
            allowed_methods=frozenset({"GET"}),
            raise_on_status=False,
        )
        self.session.mount("https://", HTTPAdapter(max_retries=retry))
        self.session.headers.update({"User-Agent": settings.user_agent})

    def _existing_download(self, task: DocumentTask) -> DownloadedDocument | None:
        if not task.chemin_local or not task.empreinte_sha256:
            return None
        candidate = (self.settings.project_root / task.chemin_local).resolve()
        if not candidate.is_relative_to(self.settings.document_root):
            return None
        if not candidate.is_file():
            return None
        detected = _detect_document_type(candidate)
        return DownloadedDocument(
            relative_path=task.chemin_local,
            absolute_path=candidate,
            sha256=task.empreinte_sha256,
            size_bytes=candidate.stat().st_size,
            mime_type=None,
            detected_type=detected,
        )

    def _download(self, task: DocumentTask) -> DownloadedDocument:
        existing = self._existing_download(task)
        if existing is not None:
            return existing

        validate_public_document_url(task.url)
        try:
            response = self.session.get(
                task.url,
                timeout=self.settings.http_timeout,
                stream=True,
                allow_redirects=True,
            )
            response.raise_for_status()
        except requests.RequestException as exc:
            raise DocumentProcessingError(
                f"téléchargement impossible: {type(exc).__name__}"
            ) from exc

        validate_public_document_url(response.url)
        declared_length = response.headers.get("Content-Length")
        if declared_length:
            try:
                if int(declared_length) > self.settings.document_max_bytes:
                    raise DocumentProcessingError("document refusé: taille déclarée excessive")
            except ValueError:
                pass

        target_dir = (self.settings.document_root / task.tender_id).resolve()
        if not target_dir.is_relative_to(self.settings.document_root):
            raise DocumentProcessingError("chemin document local invalide")
        target_dir.mkdir(parents=True, exist_ok=True)
        target = (target_dir / _safe_filename(task.url, task.id)).resolve()
        if not target.is_relative_to(target_dir):
            raise DocumentProcessingError("nom de fichier document invalide")
        temporary = target.with_suffix(target.suffix + ".part")

        digest = hashlib.sha256()
        total = 0
        try:
            with temporary.open("wb") as handle:
                for chunk in response.iter_content(chunk_size=64 * 1024):
                    if not chunk:
                        continue
                    total += len(chunk)
                    if total > self.settings.document_max_bytes:
                        raise DocumentProcessingError("document refusé: taille excessive")
                    digest.update(chunk)
                    handle.write(chunk)
            if total == 0:
                raise DocumentProcessingError("document vide")
            detected = _detect_document_type(temporary)
            os.replace(temporary, target)
        except Exception:
            temporary.unlink(missing_ok=True)
            raise

        relative = target.relative_to(self.settings.project_root).as_posix()
        return DownloadedDocument(
            relative_path=relative,
            absolute_path=target,
            sha256=digest.hexdigest(),
            size_bytes=total,
            mime_type=response.headers.get("Content-Type", "").split(";", 1)[0] or None,
            detected_type=detected,
        )

    def process(
        self,
        limit: int = 50,
        retry_errors: bool = False,
        reprocess_extracted: bool = False,
        tender_id: str | None = None,
    ) -> DocumentProcessingResult:
        if limit < 1:
            raise ValueError("limit doit être supérieur à zéro")
        self.database.check_connection()
        tasks = [
            DocumentTask.from_row(row)
            for row in self.database.documents_to_process(
                limit,
                retry_errors,
                reprocess_extracted,
                tender_id,
            )
        ]
        result = DocumentProcessingResult(demandes=len(tasks))

        try:
            for task in tasks:
                downloaded: DownloadedDocument | None = None
                try:
                    downloaded = self._download(task)
                    text = extract_text(
                        downloaded.absolute_path,
                        downloaded.detected_type,
                        self.settings,
                    )
                    if downloaded.detected_type not in SUPPORTED_DOCUMENT_TYPES:
                        status = "non_supporte"
                        result.non_supportes += 1
                    elif text:
                        status = "extrait"
                        result.extraits += 1
                    else:
                        status = "sans_texte"
                        result.sans_texte += 1

                    self.database.save_document_processing(
                        document_id=task.id,
                        relative_path=downloaded.relative_path,
                        sha256=downloaded.sha256,
                        size_bytes=downloaded.size_bytes,
                        mime_type=downloaded.mime_type,
                        detected_type=downloaded.detected_type,
                        extracted_text=text,
                        status=status,
                    )
                except Exception as exc:
                    error = self.settings.redact(exc, max_length=300)
                    self.database.fail_document_processing(
                        task.id,
                        error,
                        relative_path=(downloaded.relative_path if downloaded else None),
                        sha256=(downloaded.sha256 if downloaded else None),
                        size_bytes=(downloaded.size_bytes if downloaded else None),
                        mime_type=(downloaded.mime_type if downloaded else None),
                        detected_type=(downloaded.detected_type if downloaded else None),
                    )
                    result.echecs += 1
                    result.erreurs.append(f"document {task.id}: {error}")
            return result
        finally:
            self.session.close()
