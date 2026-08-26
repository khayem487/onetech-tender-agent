from __future__ import annotations

import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "livrables"
VISUALS = OUT_DIR / "visuals"
ARCHITECTURE = VISUALS / "architecture_v2.png"
DOCUMENT_PIPELINE = VISUALS / "pipeline_documents.png"
TRACEABILITY = VISUALS / "traceability_data.png"
AUTOMATION = VISUALS / "automation_channels.png"
RESULTS = VISUALS / "results_dashboard.png"
PPTX_OUT = OUT_DIR / "PRESENTATION_SOUTENANCE_KHAYEM_BEN_GHORBEL.pptx"
NOTES_OUT = OUT_DIR / "_presentation_notes.json"


NAVY = "0B1F3A"
BLUE = "2D65F2"
CYAN = "06A6B5"
GREEN = "28B487"
ORANGE = "F59E0B"
RED = "E35D6A"
PURPLE = "7165E3"
INK = "17243B"
MUTED = "66758A"
PALE = "F4F7FB"
WHITE = "FFFFFF"
LINE = "DCE4EF"
LIGHT_BLUE = "EAF0FF"
LIGHT_CYAN = "E4F8FA"
LIGHT_GREEN = "E8F8F2"
LIGHT_ORANGE = "FFF4DD"
LIGHT_RED = "FDECEF"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_fill(shape, color: str, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency


def set_line(shape, color: str, width: float = 1.0, transparency: int = 0) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = INK,
    bold: bool = False,
    font: str = "Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.03,
    linespacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = align
    if linespacing:
        p.line_spacing = linespacing
    for run in p.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs, x, y, w, h, *, size=18, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    p = frame.paragraphs[0]
    p.alignment = align
    p.text = ""
    for item in runs:
        run = p.add_run()
        run.text = item[0]
        run.font.name = "Aptos"
        run.font.size = Pt(item[1] if len(item) > 1 else size)
        run.font.bold = bool(item[2]) if len(item) > 2 else False
        run.font.color.rgb = rgb(item[3] if len(item) > 3 else color)
    return box


def add_rect(slide, x, y, w, h, *, fill=WHITE, line=LINE, radius=True, lw=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill)
    set_line(shape, line, lw)
    return shape


def add_circle(slide, x, y, d, *, fill=BLUE, line=None, lw=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    set_fill(shape, fill)
    if line:
        set_line(shape, line, lw)
    else:
        shape.line.fill.background()
    return shape


def add_badge(slide, text, x, y, w, *, fill=LIGHT_BLUE, color=BLUE, size=10.5):
    shape = add_rect(slide, x, y, w, 0.32, fill=fill, line=fill, radius=True, lw=0.5)
    add_text(slide, text, x + 0.05, y + 0.01, w - 0.1, 0.27, size=size, color=color, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return shape


def add_label_circle(slide, label, x, y, d=0.42, *, fill=BLUE, color=WHITE, size=12):
    add_circle(slide, x, y, d, fill=fill)
    add_text(slide, label, x, y + 0.01, d, d - 0.02, size=size, color=color, bold=True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def add_card(slide, title, body, x, y, w, h, *, accent=BLUE, fill=WHITE, icon=None,
             title_size=15, body_size=11.5):
    add_rect(slide, x, y, w, h, fill=fill, line=LINE)
    accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    set_fill(accent_bar, accent)
    accent_bar.line.fill.background()
    tx = x + 0.23
    if icon:
        add_label_circle(slide, icon, x + 0.22, y + 0.18, 0.38, fill=accent, size=10)
        tx = x + 0.72
    add_text(slide, title, tx, y + 0.17, w - (tx - x) - 0.18, 0.33, size=title_size,
             color=INK, bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, body, x + 0.23, y + 0.58, w - 0.42, h - 0.69, size=body_size,
             color=MUTED, linespacing=1.0)


def add_metric(slide, value, label, x, y, w, *, accent=BLUE, suffix=None):
    add_rect(slide, x, y, w, 1.16, fill=WHITE, line=LINE)
    add_text(slide, value, x + 0.18, y + 0.15, w - 0.35, 0.45, size=27, color=accent, bold=True)
    if suffix:
        add_text(slide, suffix, x + w - 0.58, y + 0.22, 0.38, 0.25, size=10, color=MUTED, bold=True,
                 align=PP_ALIGN.RIGHT)
    add_text(slide, label, x + 0.18, y + 0.68, w - 0.35, 0.32, size=10.5, color=MUTED, bold=True)


def add_arrow(slide, x1, y1, x2, y2, *, color=MUTED, width=1.5, dashed=False):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width)
    conn.line.end_arrowhead = True
    if dashed:
        conn.line.dash_style = 4
    return conn


def add_footer(slide, number: int, section: str):
    add_text(slide, "ONE TECH BUSINESS SOLUTIONS  ·  CY TECH", 0.68, 7.13, 5.4, 0.18,
             size=8.4, color="8A98AA", bold=True)
    add_text(slide, section.upper(), 9.1, 7.13, 3.35, 0.18, size=8.4, color="8A98AA", bold=True,
             align=PP_ALIGN.RIGHT)
    add_label_circle(slide, str(number), 12.54, 7.03, 0.28, fill=NAVY, size=8.5)


def add_header(slide, number: int, title: str, kicker: str):
    add_badge(slide, kicker.upper(), 0.68, 0.42, max(1.05, len(kicker) * 0.072 + 0.42),
              fill=LIGHT_CYAN, color=CYAN, size=9.5)
    add_text(slide, title, 0.68, 0.87, 12.0, 0.58, size=28.5, color=NAVY, bold=True,
             font="Aptos Display")
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(1.51), Inches(12.0), Inches(0.025))
    set_fill(line, LINE)
    line.line.fill.background()
    add_footer(slide, number, kicker)


def new_slide(prs, number: int, title: str, kicker: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(PALE)
    # Derive the visible number from the actual deck order so inserted slides
    # cannot desynchronise the footer numbering.
    add_header(slide, len(prs.slides), title, kicker)
    return slide


def add_bullet_list(slide, items, x, y, w, h, *, size=14, color=INK, bullet_color=CYAN,
                    gap=0.56, start=0):
    for i, item in enumerate(items):
        yy = y + i * gap
        add_label_circle(slide, str(start + i + 1) if start else "", x, yy + 0.02, 0.24,
                         fill=bullet_color, size=7.8)
        if not start:
            add_circle(slide, x + 0.07, yy + 0.12, 0.09, fill=bullet_color)
        add_text(slide, item, x + 0.38, yy, w - 0.38, min(gap, 0.52), size=size, color=color,
                 valign=MSO_ANCHOR.MIDDLE)


NOTES = []


def note(text: str) -> None:
    NOTES.append(text.strip())


def slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid(); bg.fore_color.rgb = rgb(NAVY)
    add_circle(slide, 10.10, -0.65, 4.15, fill=CYAN)
    glow = add_circle(slide, 9.15, 4.82, 2.4, fill=BLUE)
    glow.fill.transparency = 22
    for i, (xx, yy, cc) in enumerate([(10.3, 0.85, WHITE), (11.55, 2.2, GREEN), (9.45, 3.05, ORANGE), (11.3, 4.85, WHITE)]):
        add_circle(slide, xx, yy, 0.16, fill=cc)
    add_badge(slide, "SOUTENANCE DE STAGE  ·  2026", 0.82, 0.72, 2.78, fill="19365D", color="8FE2E7")
    add_text(slide, "OneTech\nTender Agent", 0.82, 1.42, 7.35, 1.65, size=43, color=WHITE, bold=True,
             font="Aptos Display")
    add_text(slide, "Système IA agentique pour collecter, analyser\net diffuser les appels d’offres publics",
             0.86, 3.25, 7.5, 0.95, size=21, color="D9E5F4")
    add_text(slide, "Khayem Ben Ghorbel", 0.86, 4.60, 4.2, 0.35, size=16, color=WHITE, bold=True)
    add_text(slide, "CY Tech  ·  One Tech Business Solutions", 0.86, 5.00, 5.8, 0.28,
             size=12.5, color="9FB1C8")
    values = [("100", "AO centralisées", CYAN), ("152", "documents traités", GREEN), ("67/67", "tests Python", ORANGE)]
    for i, (v, lab, c) in enumerate(values):
        xx = 0.86 + i * 2.25
        add_text(slide, v, xx, 6.02, 0.95, 0.42, size=23, color=c, bold=True)
        add_text(slide, lab, xx + 0.72, 6.08, 1.45, 0.30, size=9.6, color="B8C6D8", bold=True)
    add_text(slide, "01", 12.15, 6.80, 0.45, 0.25, size=9.5, color="A9BCD2", bold=True,
             align=PP_ALIGN.RIGHT)
    note("""Bonjour. Dans le cadre de mon stage chez One Tech Business Solutions, j’ai conçu OneTech Tender Agent : une chaîne complète qui collecte des appels d’offres publics, lit leurs documents, structure les faits, produit une première recommandation Go/No-Go et diffuse un digest aux équipes. Le résultat validé porte sur cent appels d’offres, cent cinquante-deux documents intégralement traités et soixante-sept tests Python réussis, complétés par six scripts SQL. Pendant cette soutenance, je vais présenter l’entreprise, la mission qui m’a été confiée, les choix techniques, les réalisations concrètes, une démonstration et les compétences que j’ai développées.""")


def slide_company(prs):
    slide = new_slide(prs, 2, "One Tech Business Solutions", "Entreprise d’accueil")
    add_rect(slide, 0.72, 1.83, 4.15, 4.82, fill=NAVY, line=NAVY)
    add_badge(slide, "ONE TECH GROUP", 1.06, 2.17, 1.78, fill="19365D", color="8FE2E7")
    add_text(slide, "Une expertise numérique\nau service des métiers", 1.06, 2.78, 3.30, 1.02,
             size=23, color=WHITE, bold=True, font="Aptos Display")
    add_text(slide, "Développement logiciel\nIntégration de systèmes\nData & intelligence artificielle\nAutomatisation des processus",
             1.06, 4.46, 3.30, 1.70, size=13.5, color="D9E6F5", linespacing=1.10)
    add_text(slide, "Mission du stage", 5.35, 1.96, 3.40, 0.36, size=16, color=CYAN, bold=True)
    add_text(slide, "Transformer la veille sur les marchés publics en un processus structuré, traçable et exploitable.",
             5.35, 2.50, 6.78, 0.95, size=23, color=NAVY, bold=True, font="Aptos Display")
    cards = [
        ("Veille", "Repérer les opportunités publiées", BLUE, "01"),
        ("Qualification", "Analyser les faits et les documents", PURPLE, "02"),
        ("Diffusion", "Transmettre un digest actionnable", GREEN, "03"),
    ]
    for i, (title, body, color, icon) in enumerate(cards):
        add_card(slide, title, body, 5.35 + i * 2.42, 4.15, 2.18, 1.55,
                 accent=color, icon=icon, title_size=13.2, body_size=10.2)
    add_badge(slide, "Stage ingénieur informatique · ING1", 7.02, 6.21, 3.54,
              fill=LIGHT_CYAN, color=CYAN)
    note("""One Tech Business Solutions est l’entreprise d’accueil du stage. Le contexte est celui des services numériques : développement logiciel, intégration de systèmes, Data, intelligence artificielle et automatisation. Dans ce cadre, la veille sur les marchés publics représente une activité utile mais très documentaire. Ma mission a consisté à transformer cette veille en un processus structuré : repérer les publications, qualifier les dossiers à partir de leurs pièces, puis transmettre une synthèse exploitable. Le projet devait donc relier une problématique métier à plusieurs disciplines informatiques, depuis la collecte web jusqu’à l’automatisation et l’IA agentique.""")


def slide_2(prs):
    slide = new_slide(prs, 2, "Du chaos documentaire à une décision vérifiable", "Problème métier")
    add_card(slide, "Information dispersée", "Listes HAICOP, pages détail et pièces jointes ne forment pas un dossier unique.",
             0.72, 1.86, 3.55, 1.28, accent=RED, fill=LIGHT_RED, icon="1")
    add_card(slide, "Formats hétérogènes", "PDF, DOCX, anciens DOC, scans et images exigent plusieurs méthodes d’extraction.",
             0.72, 3.31, 3.55, 1.28, accent=ORANGE, fill=LIGHT_ORANGE, icon="2")
    add_card(slide, "Risque opérationnel", "Échéance, pièce obligatoire, doublon ou critère éliminatoire peuvent être manqués.",
             0.72, 4.76, 3.55, 1.28, accent=RED, fill=LIGHT_RED, icon="3")
    add_arrow(slide, 4.55, 3.95, 5.65, 3.95, color=CYAN, width=2.6)
    add_rect(slide, 5.92, 1.95, 6.72, 4.65, fill=WHITE, line="BBD7DB", lw=1.4)
    add_badge(slide, "DOSSIER UNIQUE", 6.28, 2.25, 1.62, fill=LIGHT_CYAN, color=CYAN)
    add_text(slide, "Une qualification qui reste\nexplicable avant la décision", 6.28, 2.83, 5.42, 0.92,
             size=23, color=NAVY, bold=True, font="Aptos Display")
    add_bullet_list(slide, ["Faits officiels centralisés", "Opportunités automatiquement priorisées",
                            "Sources et échéances visibles", "Recommandation Go/No-Go traçable"],
                    6.30, 4.05, 5.65, 2.1, size=13.2, gap=0.52)
    add_badge(slide, "De la publication au dossier analysé", 9.13, 6.06, 2.97,
              fill=LIGHT_GREEN, color=GREEN, size=9.5)
    note("""Le point de départ est un flux d’information dispersé. Un même appel d’offres peut être décrit dans une liste, complété sur une page détaillée et accompagné de plusieurs pièces jointes. Ces pièces utilisent des formats différents et les éléments importants — échéance, organisme, critères ou documents — doivent être rapprochés. Le besoin n’était donc pas simplement de télécharger des fichiers. Il fallait construire un dossier unique, conserver la source de chaque fait, prioriser les opportunités et rendre la recommandation compréhensible. C’est cette continuité, de la publication au dossier analysé, qui structure tout le système.""")


def slide_3(prs):
    slide = new_slide(prs, 3, "Six capacités livrées de bout en bout", "Mission réalisée")
    cards = [
        ("Collecter", "Données et documents publics HAICOP", "1", BLUE),
        ("Structurer", "PostgreSQL, JSONB et anti-doublon", "2", CYAN),
        ("Extraire", "PDF, DOCX, DOC, scans et images", "3", PURPLE),
        ("Analyser", "Fiche Go/No-Go sourcée et persistée", "4", GREEN),
        ("Automatiser", "Digest quotidien Gmail et Telegram", "5", ORANGE),
        ("Interopérer", "Skills OpenClaw et serveur MCP", "6", BLUE),
    ]
    for i, (title, body, icon, color) in enumerate(cards):
        row, col = divmod(i, 3)
        add_card(slide, title, body, 0.72 + col * 4.18, 1.84 + row * 1.47, 3.83, 1.22,
                 accent=color, icon=icon, title_size=14.5, body_size=10.8)
    add_rect(slide, 0.72, 5.08, 12.0, 1.25, fill=NAVY, line=NAVY)
    labels = [("HAICOP", BLUE), ("PostgreSQL", GREEN), ("OCR", ORANGE),
              ("OpenClaw", PURPLE), ("n8n", CYAN), ("Gmail + Telegram", GREEN)]
    for i, (label, color) in enumerate(labels):
        x = 0.98 + i * 1.92
        add_circle(slide, x, 5.43, 0.22, fill=color)
        add_text(slide, label, x + 0.34, 5.35, 1.45, 0.32, size=10.3,
                 color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_badge(slide, "MVP validé sur données réelles", 4.73, 6.52, 2.88,
              fill=LIGHT_GREEN, color=GREEN)
    note("""La mission s’est matérialisée par six capacités effectivement livrées. Premièrement, la collecte des données et documents publics HAICOP. Deuxièmement, la structuration dans PostgreSQL avec anti-doublon. Troisièmement, l’extraction de texte pour les formats natifs, les scans et les anciens fichiers Word. Quatrièmement, l’analyse Go/No-Go avec sources et contexte persistés. Cinquièmement, l’automatisation du digest quotidien par n8n. Enfin, l’interopérabilité grâce aux skills OpenClaw et au serveur MCP. Ces briques ne sont pas des maquettes isolées : elles ont été reliées et validées ensemble sur les données réelles du projet.""")


def slide_4(prs):
    slide = new_slide(prs, 4, "L’agent orchestre ; les outils garantissent", "Architecture agentique")
    cx, cy = 6.05, 3.25
    add_circle(slide, cx, cy, 1.30, fill=NAVY)
    add_text(slide, "OpenClaw", cx, cy + 0.41, 1.30, 0.28, size=15, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)
    nodes = [
        (1.00, 2.00, "MODÈLE", "Raisonnement\nsémantique", PURPLE, LIGHT_BLUE),
        (9.55, 2.00, "PYTHON / SQL", "HTTP · parsing\ntransactions", BLUE, LIGHT_BLUE),
        (1.00, 4.65, "SKILLS", "Procédures\ncontrôlées", CYAN, LIGHT_CYAN),
        (9.55, 4.65, "OCR / FICHIERS", "Extraction\ndéterministe", GREEN, LIGHT_GREEN),
    ]
    for x, y, title, body, c, f in nodes:
        add_card(slide, title, body, x, y, 2.78, 1.32, accent=c, fill=f, title_size=13.5, body_size=11)
        add_arrow(slide, cx + 0.65 if x > cx else cx + 0.65, cy + 0.65,
                  x if x > cx else x + 2.78, y + 0.66, color="9AA9BA", width=1.25)
    add_rect(slide, 4.12, 5.76, 5.18, 0.74, fill=LIGHT_CYAN, line="B9E4E8")
    add_text(slide, "L’agent utilise un modèle, mais l’agent n’est pas le modèle.",
             4.35, 5.96, 4.72, 0.26, size=13.2, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_badge(slide, "Responsabilités séparées", 5.16, 1.82, 3.00, fill=LIGHT_GREEN, color=GREEN)
    note("""L’architecture distingue clairement le raisonnement et l’exécution. OpenClaw orchestre la mission : il prépare le contexte, appelle les capacités autorisées et restitue une synthèse. Le modèle intervient lorsque le sens des documents doit être interprété. Python, SQL, HTTP et l’OCR prennent en charge les opérations qui doivent rester déterministes. Les skills décrivent la procédure et limitent les actions possibles. Cette séparation rend le système plus fiable : une analyse peut évoluer sans modifier les faits collectés, et une opération technique peut être testée indépendamment du modèle.""")


def slide_5(prs):
    slide = new_slide(prs, 5, "Une chaîne de bout en bout, reproductible", "Architecture")
    if ARCHITECTURE.exists():
        slide.shapes.add_picture(str(ARCHITECTURE), Inches(0.70), Inches(1.65), Inches(12.03), Inches(5.28))
    else:
        add_text(slide, "Architecture indisponible", 4.8, 3.4, 3.5, 0.4, size=20, color=RED, bold=True)
    note("""Cette vue résume le chemin complet. HAICOP fournit les données publiques, les pages détaillées et les pièces jointes. La couche d’acquisition les récupère, contrôle les téléchargements et transforme chaque document en texte. PostgreSQL devient ensuite la source de vérité : il relie les faits, les fichiers, les analyses et les envois. OpenClaw et le profil OneTech produisent la qualification, puis n8n déclenche la diffusion par Gmail et Telegram. Le point important est la continuité : chaque étape reçoit un état vérifiable de la précédente et conserve une preuve pour la suivante.""")


def slide_6(prs):
    slide = new_slide(prs, 6, "PostgreSQL relie les faits, les preuves et la décision", "Traçabilité")
    if TRACEABILITY.exists():
        slide.shapes.add_picture(str(TRACEABILITY), Inches(0.70), Inches(1.65), Inches(12.03), Inches(5.28))
    note("""PostgreSQL est le socle de traçabilité. Les faits officiels sont conservés dans appels_offres avec l’identifiant HAICOP et le payload original en JSONB. Les preuves documentaires sont stockées séparément avec leur URL, leur empreinte SHA-256 et le texte extrait. Enfin, l’analyse conserve le hash exact du contexte, le score, la décision et la fiche. Cette organisation permet de répondre à quatre questions : quelle donnée a été reçue, quel document a été lu, sur quel contexte la décision a été produite et si le même traitement a déjà eu lieu. L’UPSERT et les contraintes uniques rendent les replays idempotents.""")


def slide_7(prs):
    slide = new_slide(prs, 7, "Du fichier brut au texte exploitable", "Documents")
    if DOCUMENT_PIPELINE.exists():
        slide.shapes.add_picture(str(DOCUMENT_PIPELINE), Inches(0.70), Inches(1.65), Inches(12.03), Inches(5.28))
    note("""Le pipeline documentaire choisit automatiquement la bonne méthode. Le fichier est d’abord téléchargé depuis un domaine public autorisé, puis son type réel et sa taille sont vérifiés. Une empreinte SHA-256 permet d’identifier le contenu et d’éviter un retraitement inutile. Les PDF texte et les DOCX utilisent une extraction native. Pour un scan ou une image, Poppler rend la page puis Tesseract réalise l’OCR en français, arabe et anglais. Les anciens DOC sont convertis temporairement avec LibreOffice avant lecture. Le résultat est de quarante-et-un PDF, soixante-quatorze DOCX, trente-cinq DOC et deux images, soit cent cinquante-deux textes persistés sur cent cinquante-deux documents.""")


def slide_8(prs):
    slide = new_slide(prs, 8, "Une recommandation liée à ce qui a été lu", "Analyse Go/No-Go")
    add_rect(slide, 0.74, 1.88, 4.20, 4.62, fill=WHITE, line=LINE)
    add_badge(slide, "CAS VALIDÉ", 1.05, 2.18, 1.32, fill=LIGHT_BLUE, color=BLUE)
    add_text(slide, "Tender-103307", 1.05, 2.72, 3.25, 0.35, size=20, color=NAVY, bold=True)
    add_text(slide, "2/2 documents lus", 1.05, 3.18, 2.40, 0.28, size=11.5, color=MUTED, bold=True)
    add_text(slide, "36", 1.05, 3.72, 1.12, 0.72, size=44, color=RED, bold=True)
    add_text(slide, "/100", 2.15, 4.08, 0.85, 0.30, size=15, color=MUTED, bold=True)
    add_badge(slide, "NO_GO", 3.18, 3.94, 1.18, fill=LIGHT_RED, color=RED, size=12)
    add_text(slide, "Analyse ID 53", 1.05, 4.75, 1.65, 0.30, size=11, color=INK, bold=True)
    add_text(slide, "Replay → même ID", 1.05, 5.14, 1.85, 0.30, size=11, color=GREEN, bold=True)
    add_text(slide, "Profil OneTech validé", 1.05, 5.53, 2.05, 0.30, size=11, color=CYAN, bold=True)
    add_text(slide, "Une décision argumentée et reproductible.", 1.05, 6.02, 3.32, 0.30,
             size=10.8, color=MUTED)
    features = [
        ("Contexte borné", "Export PostgreSQL complet", BLUE),
        ("Contexte hashé", "Décision liée au snapshot", CYAN),
        ("Fiche en 9 sections", "Sources et risques explicites", PURPLE),
        ("4 critères /25", "Adéquation · délai · admin · risque", GREEN),
        ("Validation stricte", "GO interdit si analyse partielle", ORANGE),
    ]
    for i, (title, body, c) in enumerate(features):
        add_card(slide, title, body, 5.24, 1.87 + i * 0.91, 7.45, 0.73, accent=c,
                 title_size=12.4, body_size=9.7)
    note("""L’analyse débute par un contexte préparé depuis PostgreSQL. Ce contexte contient les faits officiels, le texte des pièces et le profil métier OneTech. Il est hashé afin que la décision reste liée au snapshot réellement lu. La fiche générée suit neuf sections et attribue quatre notes sur vingt-cinq : adéquation, délai, administratif et risque. Le cas démontré, Tender-103307, a produit trente-six sur cent et une recommandation NO_GO après lecture des deux documents disponibles. L’intérêt technique du cas est surtout le replay : le même contexte retrouve l’analyse numéro cinquante-trois, sans créer de doublon et avec les mêmes preuves.""")


def slide_9(prs):
    slide = new_slide(prs, 9, "Chaque matin, le même processus — sans double envoi", "Automatisation")
    if AUTOMATION.exists():
        slide.shapes.add_picture(str(AUTOMATION), Inches(0.70), Inches(1.65), Inches(12.03), Inches(5.28))
    note("""L’automatisation quotidienne est décrite dans n8n. À huit heures, heure de Paris, les jours ouvrés, le workflow appelle une route authentifiée du worker. Celui-ci confie à OpenClaw une mission bornée, puis Python exécute la collecte, l’analyse et la génération du digest. PostgreSQL conserve l’état de diffusion par canal et par destinataire. Ce registre est essentiel : si le workflow est relancé, le même message n’est pas envoyé une seconde fois. La chaîne a été validée sur trois canaux : Mailpit pour la preuve locale, Gmail pour les destinataires et Telegram pour le digest ainsi que les réponses aux mentions du bot.""")


def slide_10(prs):
    slide = new_slide(prs, 10, "Cinq outils MCP pour réutiliser le cœur métier", "Interopérabilité")
    add_rect(slide, 0.72, 1.85, 5.86, 4.88, fill=WHITE, line=LINE)
    add_badge(slide, "MCP", 1.02, 2.13, 0.85, fill=LIGHT_BLUE, color=BLUE, size=11)
    add_text(slide, "5 outils bornés", 1.02, 2.65, 4.2, 0.40, size=22, color=NAVY, bold=True)
    tools = ["collect_tenders", "list_open_tenders", "get_tender", "analyze_tender", "generate_digest"]
    for i, tool in enumerate(tools):
        y = 3.25 + i * 0.56
        add_label_circle(slide, str(i + 1), 1.04, y, 0.32, fill=BLUE, size=8)
        add_text(slide, tool, 1.54, y + 0.02, 3.75, 0.27, size=11.5, color=INK, bold=True)
    add_text(slide, "Même logique métier · même PostgreSQL · clients compatibles",
             1.02, 6.20, 4.95, 0.28, size=10.5, color=MUTED, bold=True)
    add_rect(slide, 6.84, 1.85, 5.88, 4.88, fill=NAVY, line=NAVY)
    add_badge(slide, "GARDE-FOUS LIVRÉS", 7.18, 2.13, 2.12, fill="19365D", color="8FE2E7", size=10)
    guards = [
        ("Entrées validées", "schémas et tailles bornées", BLUE),
        ("Accès contrôlé", "aucune requête SQL arbitraire", CYAN),
        ("Source unique", "même PostgreSQL pour tous les clients", GREEN),
        ("Sorties auditables", "IDs, sources et statuts explicites", PURPLE),
    ]
    for i, (title, body, color) in enumerate(guards):
        y = 2.78 + i * 0.82
        add_label_circle(slide, str(i + 1), 7.20, y, 0.34, fill=color, size=8.5)
        add_text(slide, title, 7.72, y - 0.01, 2.05, 0.28, size=12.1, color=WHITE, bold=True)
        add_text(slide, body, 9.73, y - 0.01, 2.45, 0.34, size=10.2, color="C8D6E8")
    add_text(slide, "Une interface standard sans dupliquer la logique du projet.",
             7.18, 6.15, 4.95, 0.34, size=12.1, color="8FE2E7", bold=True,
             align=PP_ALIGN.CENTER)
    note("""Le serveur MCP rend le cœur du projet réutilisable par d’autres clients compatibles. Il expose cinq outils : collecter, lister les AO ouvertes, consulter un dossier, récupérer son analyse et générer un digest. Chaque outil possède un schéma d’entrée et des limites de taille. Il n’expose ni SQL arbitraire ni accès libre au système. Tous les clients s’appuient sur le même PostgreSQL, donc les résultats, les identifiants et les preuves restent cohérents. Pour moi, MCP a surtout été un exercice d’architecture : créer une interface standard sans recopier la logique métier déjà présente dans la CLI et le worker.""")


def slide_11(prs):
    slide = new_slide(prs, 11, "Des résultats validés sur des données réelles", "Résultats")
    if RESULTS.exists():
        slide.shapes.add_picture(str(RESULTS), Inches(0.70), Inches(1.65), Inches(12.03), Inches(5.28))
    note("""Les résultats sont mesurés à partir du snapshot de validation. Cent appels d’offres ont été centralisés. Les cent cinquante-deux documents disposent d’un fichier, d’un hash et d’un texte extrait. Les contrôles d’unicité n’ont détecté aucun doublon. La suite Python compte soixante-sept tests réussis et les six scripts SQL de validation passent également. Les cinq outils MCP ont été découverts et exécutés. Enfin, la diffusion a été validée sur Mailpit, Gmail et Telegram. Une sauvegarde complète a aussi été restaurée dans une base isolée, puis les neuf tables ont été comparées avant suppression de la base temporaire.""")


def slide_12(prs):
    slide = new_slide(prs, 12, "Une démonstration en quatre preuves", "Démo")
    steps = [
        ("01", "Données", "Afficher l’AO, ses documents et leurs hashes dans PostgreSQL", BLUE),
        ("02", "Analyse", "Lancer la qualification et ouvrir la fiche Go/No-Go sourcée", PURPLE),
        ("03", "Rejeu", "Relancer le même contexte et retrouver le même identifiant", GREEN),
        ("04", "Diffusion", "Montrer le workflow n8n, Gmail et le groupe Telegram", CYAN),
    ]
    for i, (num, title, body, color) in enumerate(steps):
        row, col = divmod(i, 2)
        x, y = 0.82 + col * 6.08, 1.93 + row * 2.05
        add_rect(slide, x, y, 5.66, 1.65, fill=WHITE, line=LINE)
        add_label_circle(slide, num, x + 0.30, y + 0.29, 0.54, fill=color, size=10)
        add_text(slide, title, x + 1.05, y + 0.25, 1.55, 0.34, size=15.5, color=NAVY, bold=True)
        add_text(slide, body, x + 1.05, y + 0.72, 4.10, 0.58, size=11.7, color=MUTED)
    add_rect(slide, 2.25, 6.18, 8.82, 0.52, fill=NAVY, line=NAVY)
    add_text(slide, "Prouver le lien entre la source, la décision, le replay et la notification",
             2.52, 6.31, 8.28, 0.24, size=11.5, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)
    note("""La démonstration suit quatre preuves et dure environ quatre minutes. Je commence par PostgreSQL pour montrer une AO, les documents reliés et leurs empreintes. Je lance ensuite la qualification afin d’ouvrir la fiche Go/No-Go et de retrouver les sources utilisées. Troisième preuve : je rejoue le même contexte et je montre que l’identifiant d’analyse reste inchangé, donc qu’aucun doublon n’est créé. Enfin, j’ouvre le workflow n8n et les notifications réellement envoyées par Gmail et Telegram. L’objectif n’est pas de montrer toutes les interfaces, mais de prouver le lien continu entre la source, la décision et la diffusion.""")


def slide_13(prs):
    slide = new_slide(prs, 13, "Des garde-fous intégrés dans chaque couche", "Qualité & sécurité")
    guards = [
        ("SHA-256", "Identifier chaque fichier et chaque contexte", BLUE, "#"),
        ("Idempotence", "Rejouer sans dupliquer données ni envois", GREEN, "↻"),
        ("Outils bornés", "Valider les entrées et limiter les actions", PURPLE, "<>"),
        ("Sources séparées", "Distinguer les faits du raisonnement généré", CYAN, "≠"),
        ("Tests", "Vérifier Python, SQL et restauration", ORANGE, "✓"),
        ("Validation humaine", "Conserver la décision commerciale finale", NAVY, "H"),
    ]
    for i, (title, body, color, icon) in enumerate(guards):
        row, col = divmod(i, 3)
        add_card(slide, title, body, 0.72 + col * 4.18, 1.84 + row * 1.62,
                 3.83, 1.36, accent=color, icon=icon, title_size=14.2, body_size=10.5)
    add_rect(slide, 1.52, 5.56, 10.30, 0.82, fill=NAVY, line=NAVY)
    add_text(slide, "Le modèle raisonne · le code vérifie · PostgreSQL conserve la preuve",
             1.85, 5.80, 9.64, 0.30, size=14.2, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)
    note("""La qualité ne repose pas sur une seule barrière. Les empreintes SHA-256 identifient les fichiers et les contextes. L’idempotence permet de rejouer une collecte, une analyse ou un envoi sans créer de doublon. Les outils valident leurs entrées et limitent les actions autorisées. Les faits officiels restent séparés du raisonnement généré. Les tests couvrent le code Python, les contraintes SQL et la restauration d’une sauvegarde. Enfin, la décision commerciale reste humaine. Ces garde-fous donnent une architecture où le modèle peut raisonner, tandis que le code vérifie et que PostgreSQL conserve la preuve de ce qui a réellement été traité.""")


def slide_14(prs):
    slide = new_slide(prs, 14, "Compétences développées pendant le stage", "Bilan personnel")
    add_rect(slide, 0.72, 1.84, 5.86, 4.78, fill=WHITE, line=LINE)
    add_badge(slide, "TECHNIQUES", 1.04, 2.15, 1.34, fill=LIGHT_BLUE, color=BLUE)
    technical = [
        ("Python & API", "collecte, parsing, worker"),
        ("PostgreSQL", "modélisation, JSONB, UPSERT"),
        ("Docker & n8n", "services et orchestration"),
        ("IA agentique", "OpenClaw, skills, MCP"),
        ("Documents", "OCR, PDF, DOCX et DOC"),
        ("Qualité", "tests, sécurité, sauvegarde"),
    ]
    for i, (title, body) in enumerate(technical):
        yy = 2.74 + i * 0.56
        add_circle(slide, 1.06, yy + 0.09, 0.12, fill=BLUE)
        add_text(slide, title, 1.36, yy, 1.55, 0.30, size=11.7, color=NAVY, bold=True)
        add_text(slide, body, 2.94, yy, 2.92, 0.30, size=10.6, color=MUTED)
    add_rect(slide, 6.84, 1.84, 5.88, 4.78, fill=NAVY, line=NAVY)
    add_badge(slide, "SOFT SKILLS", 7.16, 2.15, 1.48, fill="19365D", color="8FE2E7")
    soft = [
        ("Autonomie", "chercher, tester et documenter"),
        ("Communication", "expliquer les choix techniques"),
        ("Organisation", "avancer par jalons vérifiables"),
        ("Esprit critique", "comparer mesure et intuition"),
        ("Persévérance", "diagnostiquer les échecs réels"),
        ("Responsabilité", "protéger données et secrets"),
    ]
    for i, (title, body) in enumerate(soft):
        yy = 2.74 + i * 0.56
        add_circle(slide, 7.18, yy + 0.09, 0.12, fill=CYAN)
        add_text(slide, title, 7.48, yy, 1.60, 0.30, size=11.7, color=WHITE, bold=True)
        add_text(slide, body, 9.12, yy, 2.86, 0.30, size=10.6, color="C8D6E8")
    note("""Ce stage m’a permis de consolider des compétences techniques très variées. J’ai utilisé Python pour la collecte et les API, PostgreSQL pour la modélisation et l’idempotence, Docker et n8n pour l’infrastructure, OpenClaw et MCP pour l’architecture agentique, ainsi que Tesseract et LibreOffice pour les documents. J’ai aussi progressé sur les tests, la sécurité et la restauration. Sur le plan humain, j’ai gagné en autonomie, en organisation et en persévérance. J’ai appris à expliquer mes choix, à documenter un système complexe, à distinguer une intuition d’une mesure et à protéger les données sensibles. Ce sont ces compétences transversales qui m’ont permis d’aller jusqu’à un MVP réellement validé.""")


def slide_15(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb(NAVY)
    add_badge(slide, "CONCLUSION", 0.82, 0.62, 1.20, fill="19365D", color="8FE2E7")
    add_text(slide, "Une chaîne de décision\nexplicable, testée et réutilisable",
             0.82, 1.30, 8.9, 1.42, size=35, color=WHITE, bold=True, font="Aptos Display")
    pillars = [
        ("CENTRALISER", "faits officiels", CYAN),
        ("TRACER", "sources et contexte", BLUE),
        ("DÉDUPLIQUER", "données et envois", GREEN),
        ("AUTOMATISER", "Gmail + Telegram", ORANGE),
        ("ÉTENDRE", "skills + MCP", PURPLE),
    ]
    for i, (title, body, c) in enumerate(pillars):
        x = 0.82 + i * 2.43
        add_circle(slide, x, 3.35, 0.22, fill=c)
        add_text(slide, title, x + 0.36, 3.27, 1.80, 0.28, size=10.2, color=WHITE, bold=True)
        add_text(slide, body, x + 0.36, 3.62, 1.78, 0.28, size=9.5, color="AFC0D4")
    add_rect(slide, 0.82, 4.42, 11.72, 1.30, fill="132B4B", line="294664")
    outcomes = [("100", "AO centralisées", CYAN), ("152/152", "documents traités", GREEN),
                ("0", "doublon détecté", BLUE), ("5/5", "outils MCP", ORANGE)]
    for i, (value, label, color) in enumerate(outcomes):
        x = 1.18 + i * 2.86
        add_text(slide, value, x, 4.55, 2.30, 0.38, size=22, color=color, bold=True,
                 align=PP_ALIGN.CENTER)
        add_text(slide, label, x, 5.10, 2.30, 0.28, size=9.6, color=WHITE, bold=True,
                 align=PP_ALIGN.CENTER)
    add_text(slide, "L’IA reste encadrée par les données, le code et les contrôles métier.",
             1.45, 6.22, 10.45, 0.48, size=17, color="D9E6F5", bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Merci — Questions", 5.15, 6.83, 3.10, 0.28, size=12.5, color="8FE2E7", bold=True,
             align=PP_ALIGN.CENTER)
    note("""Pour conclure, le résultat du stage est une chaîne de décision complète et réutilisable. Les faits sont centralisés dans PostgreSQL, les documents sont transformés en texte, les analyses restent liées à leurs sources, les doublons sont bloqués et les résultats sont diffusés automatiquement. Les chiffres résument le travail livré : cent appels d’offres, cent cinquante-deux documents traités, aucun doublon détecté et cinq outils MCP opérationnels. Au-delà des technologies, le projet m’a appris à construire un système où l’intelligence artificielle reste encadrée par les données, le code et les contrôles métier. Merci pour votre attention, je suis prêt à répondre à vos questions.""")


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "OneTech Tender Agent — Soutenance de stage"
    prs.core_properties.subject = "Système IA agentique pour la gestion des appels d'offres publics"
    prs.core_properties.author = "Khayem Ben Ghorbel"
    prs.core_properties.keywords = "OneTech, OpenClaw, HAICOP, PostgreSQL, n8n, MCP, IA agentique"
    prs.core_properties.comments = "Présentation sans secrets — version finale de travail"

    for fn in [slide_1, slide_company, slide_2, slide_3, slide_4, slide_5, slide_6,
               slide_7, slide_8, slide_9, slide_10, slide_11, slide_12, slide_13,
               slide_14, slide_15]:
        fn(prs)

    # Embed the detailed speech into PowerPoint notes as well as JSON.
    for slide, speech in zip(prs.slides, NOTES):
        try:
            slide.notes_slide.notes_text_frame.text = speech
        except Exception:
            pass

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(PPTX_OUT)
    NOTES_OUT.write_text(json.dumps(NOTES, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"PPTX={PPTX_OUT}")
    print(f"NOTES={NOTES_OUT}")
    print(f"SLIDES={len(prs.slides)}")
    print(f"NOTES_COUNT={len(NOTES)}")


if __name__ == "__main__":
    build()
